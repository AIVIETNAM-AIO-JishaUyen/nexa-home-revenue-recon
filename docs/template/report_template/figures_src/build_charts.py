"""Dựng lại năm biểu đồ của báo cáo từ nguồn 1.941 bản ghi (sheet Faults).
Chạy: python3 build_charts.py
Giữ đúng tỷ lệ khung hình, bảng màu và bố cục của ảnh Excel gốc để layout LaTeX không đổi."""
import openpyxl, statistics as st
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager, ticker

WB = '/root/.claude/uploads/1f2ec24b-eb54-582c-bae3-62c9ce23239a/8c6b400e-PivotTableDashboard_1941_Agent_Handoff_v2_1.xlsx'
OUT = '/home/claude/figs'

for f in ['/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
          '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf',
          '/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf',
          '/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf']:
    font_manager.fontManager.addfont(f)
plt.rcParams.update({'font.family':'Liberation Serif','font.size':10,
                     'axes.edgecolor':'#BFBFBF','axes.linewidth':0.8,
                     'xtick.color':'#595959','ytick.color':'#595959',
                     'text.color':'#404040','axes.labelcolor':'#404040',
                     'figure.facecolor':'white','savefig.facecolor':'white'})

# ---------- đọc dữ liệu ----------
wb = openpyxl.load_workbook(WB, data_only=True)
ws = wb['Faults']; hdr=[c.value for c in ws[1]]; ix={h:i for i,h in enumerate(hdr) if h}
rows=[r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
assert len(rows)==1941, len(rows)
col=lambda n:[r[ix[n]] for r in rows]
F, S, T = col('Faults'), col('SteelType'), col('Steel_Plate_Thickness')
MN, MX = col('Minimum_of_Luminosity'), col('Maximum_of_Luminosity')

ORDER_CNT = ['Other_Faults','Bumps','K_Scatch','Z_Scratch','Pastry','Stains','Dirtiness']
def sub(f=None, s=None, v=T):
    return [v[i] for i in range(len(rows)) if (f is None or F[i]==f) and (s is None or S[i]==s)]

def save(fig, name, w, h):
    fig.set_size_inches(w/100, h/100)
    fig.savefig(f'{OUT}/{name}.png', dpi=300, bbox_inches=None)
    plt.close(fig); print('built', name)

# ================= Hình 4 — số vết lỗi theo loại lỗi và loại thép =================
fig, ax = plt.subplots()
x = range(7); wdt=0.38
a3=[len(sub(f,'A300')) for f in ORDER_CNT]; a4=[len(sub(f,'A400')) for f in ORDER_CNT]
ax.bar([i-wdt/2 for i in x], a3, wdt, color='#5B9BD5', label='Thép A300')
ax.bar([i+wdt/2 for i in x], a4, wdt, color='#ED7D31', label='Thép A400')
ax.set_title('Số vết lỗi theo loại lỗi và loại thép', fontweight='bold', fontsize=13, color='#404040', pad=12)
ax.set_ylabel('Số vết lỗi', fontweight='bold', fontsize=9)
ax.set_xticks(list(x)); ax.set_xticklabels(ORDER_CNT, rotation=45, ha='right', fontsize=9)
ax.set_ylim(0,450); ax.yaxis.set_major_locator(ticker.MultipleLocator(50))
ax.grid(axis='y', color='#D9D9D9', lw=0.6); ax.set_axisbelow(True)
for sp in ('top','right','left'): ax.spines[sp].set_visible(False)
ax.legend(frameon=False, fontsize=9, loc='center left', bbox_to_anchor=(1.01,0.5))
fig.subplots_adjust(left=0.11, right=0.79, top=0.86, bottom=0.30)
save(fig,'fig02_bang_so_luong_loi',738,373)

# ================= Hình 5 — phân tán độ sáng =================
SER = [('Pastry','#8FC1E8'),('Z_Scratch','#E8703A'),('K_Scatch','#A6A6A6'),('Stains','#FFC000'),
       ('Dirtiness','#3E7EBF'),('Bumps','#5CA83C'),('Other_Faults','#1F3E6E')]
fig, ax = plt.subplots()
for lab,c in SER:
    idx=[i for i in range(len(rows)) if F[i]==lab]
    ax.scatter([MN[i] for i in idx], [MX[i] for i in idx], s=7, c=c, linewidths=0, alpha=0.85, label=lab)
ax.set_title('Tương quan độ sáng của 7 loại lỗi', fontsize=12, color='#404040', pad=10)
ax.set_xlabel('Minimum Luminosity', fontstyle='italic', fontsize=9)
ax.set_ylabel('Maximum Luminosity', fontstyle='italic', fontsize=9)
ax.set_xlim(-8,250); ax.set_ylim(0,310)
ax.xaxis.set_major_locator(ticker.MultipleLocator(50)); ax.yaxis.set_major_locator(ticker.MultipleLocator(50))
ax.grid(color='#D9D9D9', lw=0.6); ax.set_axisbelow(True)
for sp in ('top','right','left'): ax.spines[sp].set_visible(False)
ax.legend(frameon=False, fontsize=8, ncol=7, loc='upper center', bbox_to_anchor=(0.5,-0.24),
          handletextpad=0.15, columnspacing=0.9, markerscale=2.2)
fig.subplots_adjust(left=0.11, right=0.98, top=0.88, bottom=0.30)
save(fig,'fig03_bang_scatterplot',741,375)

# ================= Hình 7 — biểu đồ hộp độ dày =================
fig, ax = plt.subplots()
bp=ax.boxplot([T], vert=True, widths=0.30, patch_artist=True, showmeans=True, whis=1.5,
              meanprops=dict(marker='x', markeredgecolor='#31859B', markersize=7),
              medianprops=dict(color='#31859B', lw=1.0),
              boxprops=dict(facecolor='#1FA087', edgecolor='#31859B', lw=1.0),
              whiskerprops=dict(color='#31859B', lw=1.0), capprops=dict(color='#31859B', lw=1.0),
              flierprops=dict(marker='o', markerfacecolor='#31859B', markeredgecolor='none', markersize=3.2))
ax.set_title('Độ dày tấm thép', fontweight='bold', fontsize=14, color='#333333', pad=12)
ax.set_ylim(0,350); ax.yaxis.set_major_locator(ticker.MultipleLocator(50))
ax.set_xticks([]); ax.grid(axis='y', color='#E8E8E8', lw=0.6); ax.set_axisbelow(True)
for sp in ('top','right'): ax.spines[sp].set_visible(False)
ax.spines['left'].set_color('#BFBFBF'); ax.spines['bottom'].set_color('#BFBFBF')
fig.subplots_adjust(left=0.13, right=0.97, top=0.85, bottom=0.08)
save(fig,'fig05_boxplot',642,357)

# ================= Hình 9 — độ dày trung bình ± SD mẫu =================
ORDER_T = sorted(ORDER_CNT, key=lambda f: st.mean(sub(f)))
means=[st.mean(sub(f)) for f in ORDER_T]; sds=[st.stdev(sub(f)) for f in ORDER_T]
cols=['#C55A11' if f=='K_Scatch' else '#1FA087' for f in ORDER_T]
fig, ax = plt.subplots()
ax.bar(range(7), means, 0.5, color=cols, yerr=sds, capsize=3.5,
       error_kw=dict(ecolor='#404040', elinewidth=0.9, capthick=0.9))
ax.set_title('Độ dày trung bình ± Độ lệch chuẩn mẫu theo loại lỗi',
             fontweight='bold', fontsize=12, color='#1F3864', pad=12)
ax.set_xticks(range(7)); ax.set_xticklabels(ORDER_T, rotation=45, ha='right', fontsize=9)
ax.set_ylim(0,200); ax.yaxis.set_major_locator(ticker.MultipleLocator(50))
ax.grid(axis='y', color='#E8E8E8', lw=0.6); ax.set_axisbelow(True)
for sp in ('top','right'): ax.spines[sp].set_visible(False)
fig.subplots_adjust(left=0.11, right=0.97, top=0.86, bottom=0.31)
save(fig,'fig07_column1',603,365)

# ================= Hình 11 — độ dày trung bình theo loại lỗi × loại thép =================
fig, ax = plt.subplots()
m3=[st.mean(sub(f,'A300')) for f in ORDER_T]; m4=[st.mean(sub(f,'A400')) for f in ORDER_T]
ax.bar([i-0.19 for i in range(7)], m3, 0.34, color='#1F3864', label='A300')
ax.bar([i+0.19 for i in range(7)], m4, 0.34, color='#FFC000', label='A400')
ax.set_title('Độ dày trung bình theo Loại lỗi và Loại thép',
             fontweight='bold', fontsize=12, color='#1F3864', pad=12)
ax.set_xticks(range(7)); ax.set_xticklabels(ORDER_T, fontsize=8.5)
ax.set_ylim(0,120); ax.yaxis.set_major_locator(ticker.MultipleLocator(20))
ax.grid(axis='y', color='#E8E8E8', lw=0.6); ax.set_axisbelow(True)
for sp in ('top','right'): ax.spines[sp].set_visible(False)
ax.legend(frameon=False, fontsize=9, loc='upper right', bbox_to_anchor=(1.0,1.16), ncol=1)
fig.subplots_adjust(left=0.09, right=0.97, top=0.83, bottom=0.13)
save(fig,'fig09_column',675,363)
