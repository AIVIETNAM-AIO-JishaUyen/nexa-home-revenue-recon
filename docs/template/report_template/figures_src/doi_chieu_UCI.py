"""Đối chiếu bảng đã làm sạch với bản gốc Steel Plates Faults trên UCI.

Đây là phép kiểm độc lập cuối cùng của Mục VI: mọi phép kiểm khác trong bài đều
dựng từ chính bộ dữ liệu đang được kiểm, nên không phát hiện được lỗi nằm trong
chính thứ sinh ra chúng. Bản gốc công bố công khai là nguồn thứ hai duy nhất.

Cách dùng
---------
1. Tải https://archive.ics.uci.edu/static/public/198/steel+plates+faults.zip
2. Giải nén, lấy Faults.NNA và Faults27x7_var
3. python3 doi_chieu_UCI.py <duong_dan_Faults.NNA> <duong_dan_workbook.xlsx>

Kết quả kỳ vọng: 0 ô lệch trên cả 27 biến, nhãn Faults và cột SteelType.
"""
import sys, hashlib
from collections import defaultdict

import openpyxl

# 34 cột của Faults.NNA: 27 biến đầu vào + 7 cột nhãn nhị phân
COLS = """X_Minimum X_Maximum Y_Minimum Y_Maximum Pixels_Areas X_Perimeter Y_Perimeter
Sum_of_Luminosity Minimum_of_Luminosity Maximum_of_Luminosity Length_of_Conveyer
TypeOfSteel_A300 TypeOfSteel_A400 Steel_Plate_Thickness Edges_Index Empty_Index
Square_Index Outside_X_Index Edges_X_Index Edges_Y_Index Outside_Global_Index
LogOfAreas Log_X_Index Log_Y_Index Orientation_Index Luminosity_Index SigmoidOfAreas
Pastry Z_Scratch K_Scatch Stains Dirtiness Bumps Other_Faults""".split()

FEATURES = COLS[:27]
LABELS = COLS[27:]
KEY = ('X_Minimum', 'X_Maximum', 'Y_Minimum', 'Y_Maximum')


def doc_uci(path):
    rows = []
    for line in open(path, encoding='utf-8'):
        parts = line.rstrip('\n').split('\t')
        if len(parts) < 34:
            continue
        rows.append(dict(zip(COLS, (float(v) for v in parts))))
    return rows


def doc_workbook(path, sheet='Faults'):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    hdr = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(hdr) if h}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    return rows, ix


def ghep(uci, rows, ix):
    """Ghép từng bản ghi theo bộ bốn cột tọa độ. Trả về danh sách cặp chỉ số."""
    theo_khoa = defaultdict(list)
    for i, u in enumerate(uci):
        theo_khoa[tuple(u[k] for k in KEY)].append(i)
    da_dung, cap, thieu = set(), [], []
    for j, r in enumerate(rows):
        khoa = tuple(float(r[ix[k]]) for k in KEY)
        ung_vien = [i for i in theo_khoa.get(khoa, []) if i not in da_dung]
        if ung_vien:
            da_dung.add(ung_vien[0])
            cap.append((j, ung_vien[0]))
        else:
            thieu.append(j)
    return cap, thieu


def main():
    nna = sys.argv[1] if len(sys.argv) > 1 else 'Faults.NNA'
    xlsx = sys.argv[2] if len(sys.argv) > 2 else 'PivotTableDashboard_1941.xlsx'

    raw = open(nna, 'rb').read()
    print(f"Nguồn      : {nna}")
    print(f"Dung lượng : {len(raw):,} byte")
    print(f"SHA-256    : {hashlib.sha256(raw).hexdigest()}")

    uci = doc_uci(nna)
    rows, ix = doc_workbook(xlsx)
    print(f"\nBản gốc   : {len(uci):,} bản ghi")
    print(f"Bảng sạch : {len(rows):,} bản ghi")

    cap, thieu = ghep(uci, rows, ix)
    print(f"Ghép được : {len(cap):,}   không ghép được: {len(thieu)}")
    if thieu:
        print("  CẢNH BÁO: có bản ghi không tìm thấy ở nguồn, dừng đối chiếu.")
        return 1

    print(f"\n{'Biến':28s} {'Số ô lệch':>10s}")
    print('-' * 40)
    tong = 0
    for c in FEATURES:
        lech = sum(1 for j, i in cap
                   if rows[j][ix[c]] is None
                   or abs(float(rows[j][ix[c]]) - uci[i][c]) > 1e-9)
        tong += lech
        print(f"{c:28s} {lech:10d}")

    nhan = 0
    for j, i in cap:
        dung = [l for l in LABELS if uci[i][l] == 1]
        if len(dung) != 1 or dung[0] != rows[j][ix['Faults']]:
            nhan += 1

    steel = 0
    if 'SteelType' in ix:
        for j, i in cap:
            mong_doi = 'A300' if uci[i]['TypeOfSteel_A300'] == 1 else 'A400'
            if mong_doi != rows[j][ix['SteelType']]:
                steel += 1

    print('-' * 40)
    print(f"{'27 biến đầu vào':28s} {tong:10d}")
    print(f"{'Nhãn Faults':28s} {nhan:10d}")
    print(f"{'Cột SteelType':28s} {steel:10d}")
    print('-' * 40)
    o = len(cap) * len(FEATURES)
    print(f"Đã đối chiếu {o:,} ô. Tổng số ô lệch: {tong + nhan + steel}.")
    return 0 if (tong + nhan + steel) == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
