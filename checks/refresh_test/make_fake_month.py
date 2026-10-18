"""make_fake_month.py – Sinh dữ liệu GIẢ của tháng 01/2026 để chạy thử làm mới file Power Query (plan refresh-test).

Chép thư mục dữ liệu sang chỗ khác, rồi trong bản sao:
- Thêm cửa hàng mới NX068 (mã POS 680, Mỹ, USD, mở ngày 05/01/2026) vào store_master.
- Thêm dòng POS tháng 01/2026 của 3 cửa hàng (NX001 – AUD, một cửa hàng EUR, NX068), có 1 hóa đơn xuất lặp (giống lỗi của tháng 3/2025).
- Thêm đơn online đặt trong tháng 01/2026 (có đơn hủy, đơn bị trả); ngày theo kiểu Mỹ MM/DD/YYYY như file thật.
- Thêm bút toán sổ cái tháng 01/2026: POS_DAILY theo cửa hàng × ngày ghi sổ, WEB_DELIVERY / WEB_RETURN theo ngày × loại tiền
  (kể cả 105 đơn đặt 12/2025 giao 01/2026 và 7 đơn trả 01/2026 đã có sẵn), 1 bút toán phí vận chuyển MJ-SHIP-2601.
- Thêm các dòng Flash tháng 01/2026, chuyển dòng TỔNG CỘNG xuống cuối.
- Đổi tên 4 file dữ liệu (giữ phần đầu tên) để thử cách "thay file": truy vấn tìm file theo phần đầu tên.
- Cập nhật bảng tổng kiểm soát QC_Summary theo dữ liệu mới, như bên cung cấp dữ liệu sẽ làm.
Ghi số mong đợi của tháng 01/2026 (tính thẳng từ dữ liệu giả, không qua Power Query) ra expected_2026_01.csv.

Chạy:  D:/Anaconda/python.exe -I checks/refresh_test/make_fake_month.py <thư mục data gốc> <thư mục data mới>
"""
import random
import shutil
import sys
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
SRC, DST = Path(sys.argv[1]), Path(sys.argv[2])
HERE = Path(__file__).resolve().parent
random.seed(2026)
PERIOD = "2026-01"
NEW_STORE = (680, "NX068", "United States", "Texas", "USD", datetime(2026, 1, 5))
NAMES = {  # tên file cũ -> tên file mới (giữ phần đầu tên)
    "POS_Sales_Export_2025.csv": "POS_Sales_Export_2025_2026-01.csv",
    "WEB_Orders_Export.csv": "WEB_Orders_Export_2026-01.csv",
    "Finance_GL_Revenue_FY2025.xlsx": "Finance_GL_Revenue_FY2025_2026-01.xlsx",
    "Sales_Flash_Report_2025.xlsx": "Sales_Flash_Report_2025_2026-01.xlsx",
}

if DST.exists():
    sys.exit(f"Thư mục {DST} đã có, hãy chọn thư mục mới")
DST.mkdir(parents=True)
for f in SRC.iterdir():
    if f.is_file():
        shutil.copy2(f, DST / NAMES.get(f.name, f.name))

# ---------------------------------------------------------------- Danh mục: lịch, tỷ giá, cửa hàng
MD = DST / "Nexa_Home_Master_Data.xlsx"
cal = pd.read_excel(MD, "calendar")
fxd = pd.read_excel(MD, "fx_daily")
fxm = pd.read_excel(MD, "fx_month_end")
sm = pd.read_excel(MD, "store_master")
work = sorted(cal.loc[cal.working_day == 1, "date"])
RATE_D = {(d, c): r for d, c, r in zip(fxd.rate_date, fxd.currency, fxd.usd_to_ccy)}
RATE_M = {c: r for p, c, r in zip(fxm.period, fxm.currency, fxm.usd_to_ccy) if p == PERIOD}


def posting(d):
    """Ngày ghi sổ = ngày làm việc đầu tiên từ ngày bán trở đi (theo lịch)."""
    return next(w for w in work if w >= d)


def jan_days(first, last):
    d = datetime(2026, 1, first)
    while d <= datetime(2026, 1, last):
        yield d
        d += timedelta(days=1)


# ---------------------------------------------------------------- POS
pos_old = pd.read_csv(DST / NAMES["POS_Sales_Export_2025.csv"], dtype=str)
eur_store = int(sm[(sm.currency == "EUR") & sm.close_date.isna()].pos_store_id.iloc[0])
cur_of = dict(zip(sm.pos_store_id, sm.currency)) | {NEW_STORE[0]: NEW_STORE[4]}
nx_of = dict(zip(sm.pos_store_id, sm.finance_store_code)) | {NEW_STORE[0]: NEW_STORE[1]}
country_of = dict(zip(sm.finance_store_code, sm.country)) | {NEW_STORE[1]: NEW_STORE[2]}
rno = int(pos_old.receipt_no.str[1:].astype(int).max())
pos_new = []
for store in [10, eur_store, NEW_STORE[0]]:
    # Bán từ 02/01 đến 30/01 (31/01 là thứ Bảy, ngày ghi sổ rơi sang tháng 2 mà lịch chưa có)
    for d in jan_days(5 if store == NEW_STORE[0] else 2, 30):
        for _ in range(2):
            rno += 1
            for line in range(1, random.randint(1, 2) + 1):
                qty = random.randint(1, 3)
                price = round(random.uniform(20, 800), 2)
                amount = round(qty * price, 2)
                disc = round(amount * random.choice([0, 0, 0.05, 0.1, 0.15]), 2)
                pos_new.append(dict(receipt_no=f"R{rno}", line_no=line, business_date=d, store_id=store,
                                    product_id=random.randint(1, 2517), qty=qty, unit_price=price, discount_amount=disc,
                                    line_amount=amount, currency=cur_of[store], exported_at=d.replace(hour=23, minute=15)))
pos_new = pd.DataFrame(pos_new)
# Một hóa đơn của NX068 ngày 14/01 bị xuất lại sáng hôm sau (dòng trùng)
dup = pos_new[(pos_new.store_id == NEW_STORE[0]) & (pos_new.business_date == datetime(2026, 1, 14))]
dup = dup[dup.receipt_no == dup.receipt_no.iloc[0]].copy()
dup["exported_at"] = datetime(2026, 1, 15, 8, 42)
pos_new = pd.concat([pos_new, dup], ignore_index=True)
pos_new["is_dup"] = [False] * (len(pos_new) - len(dup)) + [True] * len(dup)
pos_new["posting_date"] = pos_new.business_date.map(posting)
pos_new["net"] = (pos_new.line_amount - pos_new.discount_amount).round(2)
pos_new["fx"] = [RATE_D[(d, c)] for d, c in zip(pos_new.business_date, pos_new.currency)]
out = pos_new.copy()
out["business_date"] = out.business_date.dt.strftime("%Y-%m-%d")
out["exported_at"] = out.exported_at.dt.strftime("%Y-%m-%d %H:%M:%S")
for c in ["unit_price", "discount_amount", "line_amount"]:
    out[c] = out[c].map(lambda x: str(float(x)))
pos_all = pd.concat([pos_old, out[pos_old.columns].astype(str)], ignore_index=True)
pos_all.to_csv(DST / NAMES["POS_Sales_Export_2025.csv"], index=False, lineterminator="\n")

# ---------------------------------------------------------------- Online
web_old = pd.read_csv(DST / NAMES["WEB_Orders_Export.csv"], dtype=str, keep_default_na=False)
wno = int(web_old.web_order_id.str[1:].astype(int).max())
web_new = []
for d in jan_days(2, 25):
    for _ in range(4):
        wno += 1
        cur = random.choice(["AUD", "CAD", "EUR", "GBP", "USD"])
        gross = round(random.uniform(50, 6000), 2)
        disc = round(gross * random.choice([0, 0.05, 0.08, 0.12]), 2)
        web_new.append(dict(web_order_id=f"W{wno}", order_date=d, delivery_date=d + timedelta(days=random.randint(2, 5)),
                            order_status="Delivered", return_date=None, currency=cur, items_qty=random.randint(1, 8),
                            gross_amount=gross, discount_amount=disc, net_amount=round(gross - disc, 2)))
web_new = pd.DataFrame(web_new)
canc = web_new.sample(3, random_state=1).index
web_new.loc[canc, ["order_status", "delivery_date"]] = ["Cancelled", None]
ret = web_new.drop(canc)
ret = ret[ret.delivery_date <= datetime(2026, 1, 24)].sample(2, random_state=2).index
web_new.loc[ret, "order_status"] = "Returned"
web_new.loc[ret, "return_date"] = web_new.loc[ret, "delivery_date"] + timedelta(days=4)
out = web_new.copy()
for c in ["order_date", "delivery_date", "return_date"]:
    out[c] = out[c].map(lambda x: "" if pd.isna(x) else x.strftime("%m/%d/%Y"))
for c in ["gross_amount", "discount_amount", "net_amount"]:
    out[c] = out[c].map(lambda x: str(float(x)))
web_all_txt = pd.concat([web_old, out[web_old.columns].astype(str)], ignore_index=True)
web_all_txt.to_csv(DST / NAMES["WEB_Orders_Export.csv"], index=False, lineterminator="\n")
# Bảng online đủ (cũ + mới) dạng số, để sinh sổ cái và tính số mong đợi
web = web_all_txt.copy()
for c in ["order_date", "delivery_date", "return_date"]:
    web[c] = pd.to_datetime(web[c].replace("", None), format="%m/%d/%Y")
for c in ["gross_amount", "discount_amount", "net_amount"]:
    web[c] = web[c].astype(float)
web["fx"] = [RATE_D[(d, c)] for d, c in zip(web.order_date, web.currency)]

# ---------------------------------------------------------------- Sổ cái tháng 01/2026
def in_jan(s):
    return (s >= datetime(2026, 1, 1)) & (s <= datetime(2026, 1, 31))


gl_rows = []
for (sid, pd_, cur), g in pos_new[~pos_new.is_dup].groupby(["store_id", "posting_date", "currency"]):
    gl_rows.append((pd_, "5111", "POS_DAILY", nx_of[sid], "STORE", cur, -round(g.net.sum(), 2),
                    f"Doanh thu bán lẻ {nx_of[sid]} ngày {pd_:%d/%m/%Y}", None))
dl = web[(web.order_status != "Cancelled") & in_jan(web.delivery_date)]
for (d, cur), g in dl.groupby(["delivery_date", "currency"]):
    gl_rows.append((d, "5111", "WEB_DELIVERY", "WEB", "WEB", cur, -round(g.net_amount.sum(), 2), f"Doanh thu online giao ngày {d:%d/%m/%Y}", None))
rt = web[(web.order_status == "Returned") & in_jan(web.return_date)]
for (d, cur), g in rt.groupby(["return_date", "currency"]):
    gl_rows.append((d, "5213", "WEB_RETURN", "WEB", "WEB", cur, round(g.net_amount.sum(), 2), f"Hàng trả lại online ngày {d:%d/%m/%Y}", None))
SHIP = 7512.36
gl_rows.append((datetime(2026, 1, 30), "5111", "MANUAL", "WEB", "WEB", "USD", -SHIP,
                "Doanh thu phí vận chuyển tháng 01/2026 (theo bảng kê hãng vận chuyển)", "MJ-SHIP-2601"))
gl_rows.sort(key=lambda r: (r[0], r[2], r[3], r[5]))
GLF = DST / NAMES["Finance_GL_Revenue_FY2025.xlsx"]
wb = openpyxl.load_workbook(GLF)
ws = wb["GL_Revenue"]
gl_usd = []
for i, (d, acc, jt, sc, ch, cur, amt, desc, ref) in enumerate(gl_rows, 1):
    fx = RATE_M[cur]
    usd = round(amt / fx, 2)
    gl_usd.append(usd)
    ws.append([f"JE202601{i:05d}", d, PERIOD, acc, jt, sc, ch, cur, amt, fx, usd, desc, ref])
    ws.cell(ws.max_row, 2).number_format = ws.cell(2, 2).number_format
wb.save(GLF)
gl_all = pd.read_excel(GLF)

# ---------------------------------------------------------------- Flash tháng 01/2026 (theo ngày đặt / ngày bán, giá trước giảm giá, tỷ giá ngày)
MARKET = {"AUD": "Australia", "CAD": "Canada", "EUR": "Eurozone", "GBP": "United Kingdom", "USD": "United States"}
fl = []
p = pos_new.assign(market=pos_new.store_id.map(nx_of).map(country_of), usd=pos_new.line_amount / pos_new.fx)
for m, g in p.groupby("market"):
    fl.append([PERIOD, "Cửa hàng", m, g.receipt_no.nunique(), round(g.usd.sum())])
wj = web[in_jan(web.order_date)]
for cur, g in wj.groupby("currency"):
    fl.append([PERIOD, "Online", MARKET[cur], len(g), round((g.gross_amount / g.fx).sum())])
FLF = DST / NAMES["Sales_Flash_Report_2025.xlsx"]
wb = openpyxl.load_workbook(FLF)
ws = wb["Flash"]
last = ws.max_row  # dòng TỔNG CỘNG
assert ws.cell(last, 1).value == "TỔNG CỘNG"
ws.delete_rows(last)
for r in fl:
    ws.append(r)
n = ws.max_row
ws.append(["TỔNG CỘNG", None, None, f"=SUM(D6:D{n})", f"=SUM(E6:E{n})"])
wb.save(FLF)
flash = pd.read_excel(FLF, header=4)
flash = flash[flash["Tháng"].astype(str).str.match(r"^\d{4}-\d{2}$")]

# ---------------------------------------------------------------- Danh mục: cửa hàng mới, bảng tổng kiểm soát
wb = openpyxl.load_workbook(MD)
ws = wb["store_master"]
ws.append([NEW_STORE[0], NEW_STORE[1], NEW_STORE[2], NEW_STORE[3], NEW_STORE[4], NEW_STORE[5], None, None])
ws.cell(ws.max_row, 6).number_format = ws.cell(2, 6).number_format
posn = pos_all.assign(qty=pos_all.qty.astype(int), line_amount=pos_all.line_amount.astype(float),
                      discount_amount=pos_all.discount_amount.astype(float))
st = web.order_status.value_counts()
qc = [
    ("POS", len(posn)), ("POS", posn.receipt_no.nunique()),
    ("POS", f"{posn.business_date.min()} → {posn.business_date.max()}"), ("POS", int(posn.qty.sum())),
    ("POS", round(posn.line_amount.sum(), 2)), ("POS", round(posn.discount_amount.sum(), 2)),
    ("WEB", len(web)), ("WEB", f"{st['Delivered']} / {st['Cancelled']} / {st['Returned']}"), ("WEB", round(web.net_amount.sum(), 2)),
    ("GL", len(gl_all)), ("GL", round(gl_all.amount_usd.sum(), 2)),
    ("FLASH", len(flash)), ("FLASH", int(flash["Doanh số (USD)"].sum())),
]
FILE = {"POS": NAMES["POS_Sales_Export_2025.csv"], "WEB": NAMES["WEB_Orders_Export.csv"],
        "GL": NAMES["Finance_GL_Revenue_FY2025.xlsx"], "FLASH": NAMES["Sales_Flash_Report_2025.xlsx"]}
ws = wb["QC_Summary"]
for i, (k, v) in enumerate(qc):
    ws.cell(i + 2, 1).value = FILE[k]
    ws.cell(i + 2, 3).value = v
wb.save(MD)

# ---------------------------------------------------------------- Số mong đợi của tháng 01/2026 (tính thẳng, không qua Power Query)
pj = pos_new  # mọi dòng POS bán trong tháng 1 (kể cả dòng trùng)
nc = web[web.order_status != "Cancelled"]
exp = {
    "flash": flash.loc[flash["Tháng"] == PERIOD, "Doanh số (USD)"].sum(),
    "discount": -((pj.discount_amount / pj.fx).sum() + (wj.discount_amount / wj.fx).sum()),
    "pos_duplicate": -(pj[pj.is_dup].net / pj[pj.is_dup].fx).sum(),
    "cancelled": -(wj[wj.order_status == "Cancelled"].net_amount / wj[wj.order_status == "Cancelled"].fx).sum(),
    "online_timing": (nc[in_jan(nc.delivery_date)].net_amount / nc[in_jan(nc.delivery_date)].fx).sum()
                     - (nc[in_jan(nc.order_date)].net_amount / nc[in_jan(nc.order_date)].fx).sum(),
    "store_posting": 0.0, "transfer_error": 0.0,
    "returns": -(rt.net_amount / rt.fx).sum(),
    "shipping": SHIP,
    "unexplained": 0.0,
    "gl": -sum(gl_usd),
}
e = pd.DataFrame({"period": PERIOD, "item_code": list(exp), "usd": [round(v, 2) for v in exp.values()]})
e.to_csv(HERE / "expected_2026_01.csv", index=False, lineterminator="\n")
print(f"POS thêm {len(pos_new)} dòng (1 hóa đơn trùng {len(dup)} dòng), online thêm {len(web_new)} đơn, sổ cái thêm {len(gl_rows)} bút toán, Flash thêm {len(fl)} dòng")
print(e.to_string(index=False))
print("QC_Summary mới:", [v for _, v in qc])
