"""check_refresh.py – So file Power Query đã làm mới với dữ liệu giả tháng 01/2026 (plan refresh-test).

So với: số mong đợi (expected_2026_01.csv, do make_fake_month.py tính thẳng từ dữ liệu giả) và file Power Query gốc
(12 tháng năm 2025 phải giữ nguyên). Ghi kết quả ra ket_qua_2026_01.txt cạnh file này.

Chạy:  D:/Anaconda/python.exe -I checks/refresh_test/check_refresh.py <file Excel đã làm mới>
"""
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
NEW = Path(sys.argv[1])
OLD = ROOT / "excel" / "Nexa_Home_Recon_PQ.xlsx"
lines, fails = [], []


def say(s=""):
    print(s)
    lines.append(s)


def chk(name, ok, detail):
    say(f"[{'ĐẠT' if ok else 'KHÔNG ĐẠT'}] {name}: {detail}")
    if not ok:
        fails.append(name)


def rd(f, s):
    return pd.read_excel(f, s, header=2)


# 1. Bảng tổng kiểm soát
k = rd(NEW, "Kiem_tra")
chk("Kiem_tra (so với QC_Summary mới)", bool(k.dat.all()), f"{int(k.dat.sum())}/{len(k)} chỉ tiêu đạt")

# 2. Chất lượng dữ liệu: chỉ phép đếm dòng trùng được phép đổi (thêm đúng 1 hóa đơn trùng của dữ liệu giả)
q1, q0 = rd(NEW, "Chat_luong_du_lieu"), rd(OLD, "Chat_luong_du_lieu")
m = q1.merge(q0[["stt", "so_luong"]], on="stt", suffixes=("", "_goc"))
doi = m[m.so_luong != m.so_luong_goc]
chk("Chat_luong_du_lieu", list(doi.stt) == [1] and int(doi.so_luong.iloc[0] - doi.so_luong_goc.iloc[0]) == 1,
    "; ".join(f"phép {r.stt} ({r.kiem_tra}): {r.so_luong_goc} → {r.so_luong}" for r in doi.itertuples()) or "không phép nào đổi")

# 3. Bảng cầu nối
b1, b0 = rd(NEW, "BR_Long"), rd(OLD, "BR_Long")
per = sorted(b1.period.unique())
chk("Số tháng của bảng cầu nối", len(per) == 13 and per[-1] == "2026-01", f"{len(per)} tháng, {per[0]} → {per[-1]}; {len(b1)} dòng")
ck = b1[b1.item_code == "check"]
chk("Dòng kiểm tra (Flash + các khoản − sổ cái)", bool((ck.usd.abs() <= 1).all()), f"{len(ck)} bảng × tháng, lớn nhất {ck.usd.abs().max():.6f} USD")
mm = b0.merge(b1, on=["bridge", "period", "item_code"], how="left", suffixes=("_goc", "_moi"))
chk("12 tháng năm 2025 giữ nguyên", bool(mm.usd_moi.notna().all() and (mm.usd_goc - mm.usd_moi).abs().max() < 0.005),
    f"{len(mm)} dòng, lệch lớn nhất {(mm.usd_goc - mm.usd_moi).abs().max():.6f} USD")
e = pd.read_csv(HERE / "expected_2026_01.csv")
j = b1[(b1.bridge == "COMPANY") & (b1.period == "2026-01")].set_index("item_code").usd
e["power_query"] = e.item_code.map(j).fillna(0).round(2)
e["lech"] = (e.power_query - e.usd).round(2)
chk("Tháng 01/2026 so với số mong đợi", bool((e.lech.abs() <= 0.01).all()), f"{len(e)} khoản, lệch lớn nhất {e.lech.abs().max():.2f} USD")
say(e.rename(columns={"usd": "mong_doi"}).to_string(index=False))
say("Các khoản còn lại của tháng 01/2026 (không có số mong đợi riêng): "
    + "; ".join(f"{c} = {j[c]:,.2f}" for c in ["fx", "rounding", "check"]))
cb = rd(NEW, "Canh_bao_cau_noi")
canh = cb[cb.trang_thai != "Đạt"]
say(f"Cảnh báo cầu nối: {len(cb)} bảng × tháng, {len(canh)} cần xem: "
    + "; ".join(f"{r.bridge} {r.period} ({r.ghi_chu})" for r in canh.itertuples()))

# 4. Đối soát chi tiết của tháng 01/2026
ch = rd(NEW, "Doi_soat_CH_ngay")
c1 = ch[ch.period == "2026-01"]
chk("Cửa hàng × ngày ghi sổ, tháng 01/2026", bool((c1.match_status == "Khớp").all()) and "NX068" in set(c1.store_code),
    f"{len(c1)} ô, {int((c1.match_status == 'Khớp').sum())} khớp; NX068 có {int((c1.store_code == 'NX068').sum())} ô")
for s, ten in [("Doi_soat_online", "Online × ngày giao × loại tiền"), ("Doi_soat_hang_tra", "Hàng trả × ngày trả × loại tiền")]:
    d = rd(NEW, s)
    d1 = d[d.period == "2026-01"]
    chk(f"{ten}, tháng 01/2026", bool((d1.match_status == "Khớp").all()), f"{len(d1)} ô, {int((d1.match_status == 'Khớp').sum())} khớp")

# 5. Nhật ký ngoại lệ: tổng cột tiền = chênh lệch của cả kỳ (nay 13 tháng)
nk = rd(NEW, "Nhat_ky_du_lieu")
co = b1[b1.bridge == "COMPANY"]
delta = co[co.item_code == "gl"].usd.sum() - co[co.item_code == "flash"].usd.sum()
chk("Nhật ký: tổng = chênh lệch cả kỳ", abs(nk["Ảnh hưởng (USD)"].sum() - delta) < 0.01,
    f"{len(nk)} dòng, tổng {nk['Ảnh hưởng (USD)'].sum():,.2f}; sổ cái − Flash 13 tháng {delta:,.2f}")

# 6. Sheet file mẫu: chỉ có 12 cột tháng của năm 2025, dòng kiểm tra vẫn bằng 0
wb = openpyxl.load_workbook(NEW, data_only=True)
ws = wb["Cau_noi"]
row = next(r for r in ws.iter_rows() if str(r[0].value or "").startswith("Kiểm tra"))
vals = [c.value for c in row[1:14]]
chk("File mẫu Cau_noi: dòng kiểm tra", all(abs(v or 0) <= 0.005 for v in vals), f"cột tháng: {[c.value for c in ws[1][1:13]][:2]} … (12 cột)")

say()
say("KẾT LUẬN: " + ("mọi phép so đạt" if not fails else "không đạt: " + ", ".join(fails)))
open(HERE / "ket_qua_2026_01.txt", "w", encoding="utf-8", newline="\n").write("\n".join(lines) + "\n")
