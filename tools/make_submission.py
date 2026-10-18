"""make_submission.py – Gom bộ file nộp vào thư mục submission/ từ các file đã dựng trong dự án.

Chạy (từ thư mục gốc dự án), sau khi đã dựng lại file Power Query, làm mới Power BI và chuyển báo cáo ra PDF:
    D:/Anaconda/python.exe -I tools/make_submission.py

Tạo ra:
- Bao_cao_Nexa_Home_2025.pdf          báo cáo (report/main.pdf)
- Nexa_Home_Recon_PQ.xlsx              file Power Query (có sheet Cau_noi, Nhat_ky_ngoai_le theo file mẫu, đọc bằng công thức)
- Mau_nop_Bang_cau_noi.xlsx            file mẫu nộp của đề, điền GIÁ TRỊ lấy từ file Power Query (không còn công thức trỏ sang file khác)
- NexaRecon_PowerBI.zip                dự án Power BI dạng thư mục (.pbip) và các truy vấn kiểm tra
- Ma_Power_Query.zip                   code các truy vấn (.m) và script dựng / làm mới file Excel
- README.md                            hướng dẫn đọc bộ file
"""
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission"
PQ = ROOT / "excel" / "Nexa_Home_Recon_PQ.xlsx"

if OUT.exists():
    shutil.rmtree(OUT)
OUT.mkdir()

# 1. Báo cáo và file Power Query
shutil.copy2(ROOT / "report" / "main.pdf", OUT / "Bao_cao_Nexa_Home_2025.pdf")
shutil.copy2(PQ, OUT / "Nexa_Home_Recon_PQ.xlsx")

# 2. File mẫu nộp: chép giá trị từ sheet Cau_noi, Nhat_ky_du_lieu của file Power Query vào bản sao file mẫu của đề
pq = openpyxl.load_workbook(PQ, data_only=True)
mau = openpyxl.load_workbook(ROOT / "data" / "Mau_nop_Bang_cau_noi.xlsx")
src, dst = pq["Cau_noi"], mau["Cau_noi"]
for r in range(2, 16):
    if r == 14:
        continue                                      # dòng 14 giữ công thức cộng của file mẫu
    if r <= 13:
        dst.cell(r, 1).value = src.cell(r, 1).value   # tên khoản ("Khoản giải thích 1: Chiết khấu" ...)
    for c in range(2, 14):                            # cột B..M = 12 tháng; cột N (cả năm) giữ công thức của file mẫu
        v = src.cell(r, c).value
        dst.cell(r, c).value = v                      # giữ đủ số lẻ để dòng Kiểm tra của file mẫu ra đúng 0
nk = pd.read_excel(PQ, "Nhat_ky_du_lieu", header=2)
ws = mau["Nhat_ky_ngoai_le"]
ghi_chu = ws.cell(4, 1).value                         # dòng ghi chú của file mẫu, đưa xuống cuối
for r in range(2, 5):
    for c in range(1, 11):
        ws.cell(r, c).value = None
for i, row in enumerate(nk.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(i, c).value = round(float(v), 2) if c == 5 else v
ws.cell(len(nk) + 3, 1).value = ghi_chu
mau.save(OUT / "Mau_nop_Bang_cau_noi.xlsx")

# Kiểm tra lại: dòng kiểm tra (Flash + các khoản − sổ cái) bằng 0 ở 12 tháng; tổng nhật ký = sổ cái − Flash cả năm
chk = openpyxl.load_workbook(OUT / "Mau_nop_Bang_cau_noi.xlsx")["Cau_noi"]
lech = [round(sum(chk.cell(r, c).value or 0 for r in range(2, 14)) - (chk.cell(15, c).value or 0), 2) for c in range(2, 14)]
flash = sum(chk.cell(2, c).value for c in range(2, 14))
gl = sum(chk.cell(15, c).value for c in range(2, 14))
print("File mẫu nộp: dòng kiểm tra 12 tháng =", lech, "| tổng nhật ký", round(nk["Ảnh hưởng (USD)"].sum(), 2), "| sổ cái − Flash", round(gl - flash, 2))
assert all(abs(x) <= 0.01 for x in lech) and abs(nk["Ảnh hưởng (USD)"].sum() - (gl - flash)) < 0.01


# 3. Dự án Power BI và code Power Query, nén zip (bỏ file riêng của từng máy)
def nen(zip_name, files):
    with zipfile.ZipFile(OUT / zip_name, "w", zipfile.ZIP_DEFLATED) as z:
        for f in files:
            if f.is_file() and f.name not in ("localSettings.json", "cache.abf") and "__pycache__" not in f.parts:
                z.write(f, f.relative_to(ROOT))
    print("ghi", zip_name)


pb = ROOT / "powerbi"
nen("NexaRecon_PowerBI.zip", [pb / "NexaRecon.pbip", pb / "README.md", pb / "dax_measures.dax", pb / "build_report.py", pb / "dax_query.ps1"]
    + sorted((pb / "NexaRecon.SemanticModel").rglob("*")) + sorted((pb / "NexaRecon.Report").rglob("*")) + sorted((pb / "checks").rglob("*")))
ex = ROOT / "excel"
nen("Ma_Power_Query.zip", sorted((ex / "queries").glob("*")) + [ex / "build_workbook.ps1", ex / "refresh_workbook.ps1", ex / "huong_dan.txt"])

# 4. Hướng dẫn đọc bộ file
(OUT / "README.md").write_text("""# Bộ file nộp – Project 1: Đối soát doanh thu Nexa Home 2025

| File | Là gì | Mở thế nào |
|---|---|---|
| `Bao_cao_Nexa_Home_2025.pdf` | Báo cáo 8 phần và phụ lục. Đọc hộp tóm tắt ở trang 1 trước | Trình đọc PDF |
| `Nexa_Home_Recon_PQ.xlsx` | File Power Query: nạp 4 file dữ liệu, đối soát, bảng cầu nối, nhật ký ngoại lệ, kiểm tra chất lượng dữ liệu. Sheet `Huong_dan` liệt kê mọi sheet | Excel 365. Sửa ô `Cai_dat!B2` thành thư mục chứa dữ liệu của đề rồi chọn Data > Refresh All (khoảng 3–5 phút) |
| `Mau_nop_Bang_cau_noi.xlsx` | File mẫu nộp của đề, đã điền số (giá trị chép từ file Power Query; dòng Kiểm tra bằng 0 ở 12 tháng) | Excel |
| `NexaRecon_PowerBI.zip` | Dự án Power BI dạng thư mục: mô hình, dashboard 4 trang, toàn bộ công thức DAX (`powerbi/dax_measures.dax`), truy vấn kiểm tra và ảnh chụp (`powerbi/checks/`) | Giải nén, mở `powerbi/NexaRecon.pbip` bằng Power BI Desktop; sửa 2 tham số `DataFolder`, `PQFile` (Home > Transform data > Edit parameters) rồi Refresh. Xem `powerbi/README.md` |
| `Ma_Power_Query.zip` | Code các truy vấn Power Query (file `.m`) và script dựng lại / làm mới file Excel | Trình soạn thảo chữ |

Chênh lệch = Sổ cái − Flash; số âm nghĩa là sổ cái thấp hơn Flash. Đơn vị USD.
""", encoding="utf-8")
print("Xong:", ", ".join(sorted(p.name for p in OUT.iterdir())))
